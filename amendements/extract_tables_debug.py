from docx import Document

# Fonction récursive pour extraire et afficher tous les tableaux imbriqués
def print_tables(table, level=0):
    prefix = '  ' * level
    print(f"{prefix}Tableau (niveau {level}) :")
    for i, row in enumerate(table.rows):
        row_content = []
        for j, cell in enumerate(row.cells):
            # Vérifier s'il y a un ou plusieurs sous-tableaux dans la cellule
            if cell.tables:
                row_content.append(f"[Sous-tableau dans cellule {j}]")
            else:
                # Afficher le texte de la cellule
                row_content.append(cell.text.strip().replace('\n', ' | '))
        print(f"{prefix}  Ligne {i}: {row_content}")
        # Afficher les sous-tableaux de chaque cellule
        for cell in row.cells:
            for subtable in cell.tables:
                print_tables(subtable, level=level+1)

def is_nested_table(table, doc):
    # Vérifie si le tableau est un sous-tableau d'une cellule
    for parent_table in doc.tables:
        for row in parent_table.rows:
            for cell in row.cells:
                if table in cell.tables:
                    return True
    return False

if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python extract_tables_debug.py chemin/vers/fichier.docx")
        sys.exit(1)
    docx_path = sys.argv[1]
    doc = Document(docx_path)
    for idx, table in enumerate(doc.tables):
        if not is_nested_table(table, doc):
            print(f"Tableau principal {idx} :")
            print_tables(table) 
