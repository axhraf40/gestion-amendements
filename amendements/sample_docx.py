"""
Génère des documents d'amendements FICTIFS au même format que les vrais fichiers Word
(en-tête + tableau « النص كما جاء في المشروع / نص التعديل / التعليل »).

Utilisé par les tests et par scripts/generate_sample_docx.py.
Tout le contenu est inventé : aucun document réel n'est publié dans le dépôt.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

GROUPE = "9"

# Chaque amendement : ce qui est écrit dans le document + les valeurs attendues en base.
AMENDEMENTS = [
    {
        "entete": ["التعديل رقم : 1", f"الفريق {GROUPE}", "إضافة مادة جديدة"],
        "lignes": ["مدونة الجمارك والضرائب غير المباشرة", "المادة 3", "البند I", "الفصل 70 المكرر"],
        "titre_original": "النص الأصلي",
        "original": ["الفصل 70 المكرر – نص تجريبي أصلي للفصل.", "فقرة ثانية من النص الأصلي."],
        "modifie": ["الفصل 70 المكرر – نص تجريبي بعد التعديل.", "فقرة ثانية معدلة."],
        "justification": ["تعليل تجريبي للتعديل رقم 1."],
        "attendu": ("1", GROUPE, "إضافة مادة جديدة", "مدونة الجمارك والضرائب غير المباشرة",
                    "3", "I", "70 المكرر"),
    },
    {
        "entete": ["التعديل رقم : 2", f"الفريق {GROUPE}", ""],
        "lignes": ["الضرائب الداخلية على الاستهلاك", "المادة 5", "البند II", "الفصل 9"],
        "titre_original": "النص كما جاء في المشروع",
        # texte / tableau / texte : l'ordre doit être conservé
        "original": ["الفصل 9 – تحدد المقادير وفق الجدول التالي :",
                     ("tableau", [["بيان المنتجات", "وحدة التحصيل", "المقدار"],
                                  ["منتج تجريبي أ", "100 كلغ", "120"],
                                  ["منتج تجريبي ب", "100 كلغ", "80"]]),
                     "(الباقي لا تغيير فيه)"],
        "modifie": ["الفصل 9 – تحدد المقادير الجديدة وفق الجدول التالي :",
                    ("tableau", [["بيان المنتجات", "وحدة التحصيل", "المقدار"],
                                 ["منتج تجريبي أ", "100 كلغ", "90"]]),
                    "(الباقي لا تغيير فيه)"],
        "justification": ["تعليل تجريبي للتعديل رقم 2."],
        "attendu": ("2", GROUPE, "", "الضرائب الداخلية على الاستهلاك", "5", "II", "9"),
    },
    {
        "entete": ["التعديل رقم : ١٢", f"الفريق {GROUPE}", "إضافة مادة جديدة"],
        "lignes": ["المدونة العامة للضرائب", "المادة 8", "البند I", "المادة 42 المكرر مرّتين"],
        "titre_original": "النص كما جاء في المشروع",
        "original": ["المادة 42 المكررة مرتين : شروط التطبيق (نص تجريبي)."],
        "modifie": ["المادة 42 المكررة مرتين : شروط التطبيق بعد التعديل (نص تجريبي)."],
        "justification": ["تعليل تجريبي للتعديل رقم 12."],
        "attendu": ("12", GROUPE, "إضافة مادة جديدة", "المدونة العامة للضرائب", "8", "I",
                    "42 المكرر مرتين"),
    },
    {
        "entete": ["التعديل رقم : 13", f"الفريق {GROUPE}", ""],
        "lignes": ["جدول التعريفة الجمركية"],
        "titre_original": "النص كما جاء في المشروع",
        "original": [("tableau", [["الرسم", "الوحدة"], ["0409.00 منتج تجريبي", "10"]])],
        "modifie": [("tableau", [["الرسم", "الوحدة"], ["0409.00 منتج تجريبي", "2,5"]])],
        "justification": ["تعليل تجريبي للتعديل رقم 13."],
        "attendu": ("13", GROUPE, "", "جدول التعريفة الجمركية", "", "", ""),
    },
    {
        "entete": ["التعديل رقم : 14", f"الفريق {GROUPE}", "حذف"],
        "lignes": ["المدونة العامة للضرائب", "المادة 8", "البند III", "المادة 6"],
        "titre_original": "النص الأصلي",
        "original": ["المادة 6 – نص تجريبي سيتم حذفه."],
        "modifie": ["يحذف."],
        "justification": ["تعليل تجريبي للتعديل رقم 14."],
        "attendu": ("14", GROUPE, "حذف", "المدونة العامة للضرائب", "8", "III", "6"),
    },
]


def _rtl(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return paragraph


def _remplir_cellule(cell, contenu):
    """contenu : liste de paragraphes (str) et/ou ("tableau", lignes) pour un tableau imbriqué."""
    cell.text = ""
    premier = True
    for bloc in contenu:
        if isinstance(bloc, tuple):
            lignes = bloc[1]
            sous = cell.add_table(rows=len(lignes), cols=len(lignes[0]))
            for i, ligne in enumerate(lignes):
                for j, valeur in enumerate(ligne):
                    sous.cell(i, j).text = valeur
            premier = False
            continue
        p = cell.paragraphs[0] if premier else cell.add_paragraph()
        p.text = bloc
        _rtl(p)
        premier = False


def construire_document(amendements=AMENDEMENTS, groupe=GROUPE):
    doc = Document()
    _rtl(doc.add_paragraph(f"الفريق {groupe}"))
    _rtl(doc.add_paragraph("التعديلات المقترحة"))
    for a in amendements:
        entete = doc.add_table(rows=1, cols=3)
        for i, valeur in enumerate(a["entete"]):
            entete.cell(0, i).text = valeur
        for ligne in a["lignes"]:
            _rtl(doc.add_paragraph(ligne))
        tableau = doc.add_table(rows=2, cols=3)
        for i, titre in enumerate([a["titre_original"], "نص التعديل", "التعليل"]):
            tableau.cell(0, i).text = titre
        _remplir_cellule(tableau.cell(1, 0), a["original"])
        _remplir_cellule(tableau.cell(1, 1), a["modifie"])
        _remplir_cellule(tableau.cell(1, 2), a["justification"])
        doc.add_paragraph("")
    return doc


def enregistrer(chemin, amendements=AMENDEMENTS):
    construire_document(amendements).save(chemin)
    return chemin
