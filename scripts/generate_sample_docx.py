"""Generate a fictional amendments document to try the app.

Usage:
    python scripts/generate_sample_docx.py [output.docx]

All content is invented for demonstration purposes.
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

RED = RGBColor(0xC0, 0x00, 0x00)
BLUE = RGBColor(0x1F, 0x4E, 0x79)

AMENDMENTS = [
    {
        'numero': '1', 'type': 'تعديل', 'loi': 'المدونة العامة للضرائب',
        'article': 'المادة 6', 'alinea': 'البند I',
        'original': 'تعفى من الضريبة على الشركات الجمعيات التي لا تهدف إلى الربح.',
        'ajout': 'والتعاونيات الصغيرة التي لا يتجاوز رقم معاملاتها مليون درهم.',
        'justification': 'تشجيع الاقتصاد الاجتماعي والتضامني ودعم المقاولات الصغيرة.',
    },
    {
        'numero': '2', 'type': 'إضافة مادة جديدة', 'loi': 'مدونة الجمارك والضرائب غير المباشرة',
        'article': 'المادة 3', 'alinea': 'البند II',
        'original': 'تطبق الرسوم الجمركية وفق جدول التعريفة الجاري به العمل.',
        'ajout': 'مع تخفيض نسبة الرسوم على المعدات الطبية إلى 2,5%.',
        'justification': 'تخفيف كلفة التجهيزات الصحية وتحسين الولوج إلى العلاج.',
    },
    {
        'numero': '3', 'type': 'حذف', 'loi': 'المدونة العامة للضرائب',
        'article': 'المادة 8', 'alinea': 'البند III',
        'original': 'تفرض ضريبة إضافية على عمليات التحويل الإلكتروني.',
        'ajout': '',
        'justification': 'تشجيع الأداء الإلكتروني والحد من التعامل النقدي.',
    },
]


def rtl_paragraph(cell_or_doc, text='', bold=False, color=None, size=None,
                  align=WD_ALIGN_PARAGRAPH.RIGHT, highlight=None, strike=False):
    paragraph = cell_or_doc.add_paragraph()
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.strike = strike
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    if highlight:
        run.font.highlight_color = highlight
    return paragraph


def fill(cell, text, **kwargs):
    cell.text = ''
    first = cell.paragraphs[0]
    first.alignment = kwargs.pop('align', WD_ALIGN_PARAGRAPH.RIGHT)
    run = first.add_run(text)
    run.bold = kwargs.get('bold', False)
    if kwargs.get('color'):
        run.font.color.rgb = kwargs['color']
    if kwargs.get('highlight'):
        run.font.highlight_color = kwargs['highlight']
    if kwargs.get('strike'):
        run.font.strike = True
    return run


def build(path):
    doc = Document()
    rtl_paragraph(doc, 'الفريق النموذجي', bold=True, size=16, color=BLUE,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    rtl_paragraph(doc, 'التعديلات المقترحة على مشروع قانون المالية', bold=True, size=14,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    rtl_paragraph(doc, '(وثيقة تجريبية - محتوى غير حقيقي)', size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    for amendment in AMENDMENTS:
        table = doc.add_table(rows=4, cols=3)
        table.style = 'Table Grid'
        fill(table.cell(0, 0), f"التعديل رقم : {amendment['numero']}", bold=True, color=BLUE)
        fill(table.cell(0, 1), 'الفريق النموذجي', bold=True)
        fill(table.cell(0, 2), amendment['type'], bold=True, highlight=WD_COLOR_INDEX.YELLOW)

        fill(table.cell(1, 0), amendment['loi'], bold=True)
        fill(table.cell(1, 1), amendment['article'])
        fill(table.cell(1, 2), amendment['alinea'])

        merged = table.cell(2, 0).merge(table.cell(2, 2))
        fill(merged, 'النص الأصلي : ', bold=True)
        run = merged.paragraphs[0].add_run(amendment['original'])
        run.font.strike = amendment['type'] == 'حذف'
        if amendment['ajout']:
            p = merged.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            label = p.add_run('النص المقترح : ')
            label.bold = True
            added = p.add_run(amendment['ajout'])
            added.font.color.rgb = RED
            added.bold = True

        merged = table.cell(3, 0).merge(table.cell(3, 2))
        fill(merged, 'التعليل : ', bold=True)
        merged.paragraphs[0].add_run(amendment['justification']).italic = True

        doc.add_paragraph()

    doc.save(path)
    return path


if __name__ == '__main__':
    output = sys.argv[1] if len(sys.argv) > 1 else 'exemple_amendements.docx'
    print(f'Created {build(output)}')
